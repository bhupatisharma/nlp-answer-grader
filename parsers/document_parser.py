"""Extract selectable text from supported document formats."""

from pathlib import Path

from docx import Document
from pypdf import PdfReader


class DocumentParseError(ValueError):
    """Raised when a document cannot be read or contains no usable text."""


def extract_text(path: str | Path) -> str:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".txt":
            text = file_path.read_text(encoding="utf-8-sig")
        elif suffix == ".docx":
            document = Document(file_path)
            blocks = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                blocks.extend(" ".join(cell.text for cell in row.cells) for row in table.rows)
            text = "\n".join(block for block in blocks if block.strip())
        elif suffix == ".pdf":
            reader = PdfReader(str(file_path))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            raise DocumentParseError("Supported formats are PDF, DOCX, and TXT.")
    except DocumentParseError:
        raise
    except Exception as error:
        raise DocumentParseError(f"Could not read this {suffix.lstrip('.').upper()} file: {error}") from error

    cleaned = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not cleaned:
        hint = " Scanned PDFs need OCR; this demo reads selectable PDF text." if suffix == ".pdf" else ""
        raise DocumentParseError(f"No selectable text was found in the document.{hint}")
    return cleaned